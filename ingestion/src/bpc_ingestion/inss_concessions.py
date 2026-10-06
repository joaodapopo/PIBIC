"""Monthly public concession counts, no person-level persistence or CNJ matching."""
import hashlib
import io
import re
import uuid
from collections import Counter
from datetime import datetime, timezone

from openpyxl import load_workbook
from sqlalchemy import select

from .inss_benefits import UF_NAMES, fold
from .models import IndicadorInssConcessao

DATASET = 'beneficios-concedidos-plano-de-dados-abertos-jun-2023-a-jun-2025'


def numeric_code(value):
    if isinstance(value, bool):
        raise ValueError('Código booleano não é código INSS')
    text = str(value if value is not None else '').strip()
    if not re.fullmatch(r'\d+', text):
        raise ValueError('Código INSS não numérico; não inferir pela descrição')
    return int(text)


def aggregate_concession_rows(rows, uf, competence):
    if uf.upper() not in UF_NAMES:
        raise ValueError('UF brasileira inválida')
    for line in range(1, 21):
        header = next(rows, None)
        if header is None:
            raise ValueError('Concessões INSS sem cabeçalho')
        columns = {}
        for index, value in enumerate(header):
            columns.setdefault(re.sub(r'[^a-z0-9]', '', fold(value)), []).append(index)
        if 'especie' in columns and 'despacho' in columns:
            break
    for key in ('especie', 'despacho'):
        indices = columns.get(key, [])
        if len(indices) != 2 or indices[1] != indices[0] + 1:
            raise ValueError(f'Concessões: {key} deve ter par código/descrição')
    for key in ('competenciaconcessao', 'uf'):
        if len(columns.get(key, [])) != 1:
            raise ValueError(f'Concessões: coluna {key} ausente/ambígua')
    period_index, state_index = columns['competenciaconcessao'][0], columns['uf'][0]
    species_index, dispatch_index = columns['especie'][0], columns['despacho'][0]
    minimum = max(period_index, state_index, species_index, dispatch_index + 1)
    counts = Counter()
    read = unknown = 0
    for row in rows:
        if not any(value is not None for value in row):
            continue
        read += 1
        if len(row) <= minimum:
            raise ValueError(f'Concessões: linha {read+line} incompleta')
        period = row[period_index]
        period = period.strftime('%Y%m') if isinstance(period, datetime) else str(period).strip()
        if period != str(competence):
            raise ValueError(f'Concessões: competência divergente na linha {read+line}')
        species, dispatch = numeric_code(row[species_index]), numeric_code(row[dispatch_index])
        if species not in (87, 88):
            continue
        state = fold(row[state_index]).strip()
        if state.upper() not in UF_NAMES and state not in UF_NAMES.values():
            unknown += 1
        if state not in (uf.lower(), UF_NAMES[uf.upper()]):
            continue
        label = str(row[dispatch_index+1] or '').strip() or 'Não informado'
        counts[(uf.upper(), species, dispatch, label)] += 1
    return counts, {'linhas_lidas': read, 'linhas_selecionadas': sum(counts.values()),
                    'bpc_uf_desconhecida': unknown, 'linha_cabecalho': line,
                    'cabecalho': list(header), 'uf': uf.upper(), 'especies': [87, 88]}


def aggregate_concession_xlsx(body, uf, competence):
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Concessões INSS com múltiplas abas não suportadas')
        return aggregate_concession_rows(iter(workbook.active.iter_rows(values_only=True)), uf, competence)
    finally:
        workbook.close()


def store_concessions(session, counts, resource, competence, body, raw, run):
    digest, inserted = hashlib.sha256(body).hexdigest(), 0
    for (uf, species, code, label), quantity in counts.items():
        key = dict(fonte='inss_concessoes', competencia=competence, uf=uf, especie=species,
                   codigo_despacho=code, despacho=label, hash_arquivo=digest)
        existing = session.scalar(select(IndicadorInssConcessao).filter_by(**key))
        if existing:
            if existing.quantidade != quantity:
                raise ValueError('Concessões: mesmos bytes produziram agregação divergente')
            continue
        session.add(IndicadorInssConcessao(**key, quantidade=quantity, recurso_id=str(resource['id']),
                    recurso_url=resource['url'], arquivo_bruto=str(raw), coleta_id=uuid.UUID(run),
                    coletado_em=datetime.now(timezone.utc)))
        inserted += 1
    return inserted
