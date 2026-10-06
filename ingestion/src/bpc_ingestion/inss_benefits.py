"""Agrega indeferimentos públicos do INSS sem persistir microdados na Silver."""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from .models import IndicadorInssIndeferimento
from .public_sources import StjClient


MONTHS = ("janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto",
          "setembro", "outubro", "novembro", "dezembro")
UF_NAMES = dict(zip(
    "AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split(),
    "acre;alagoas;amapa;amazonas;bahia;ceara;distrito federal;espirito santo;goias;maranhao;mato grosso;mato grosso do sul;minas gerais;para;paraiba;parana;pernambuco;piaui;rio de janeiro;rio grande do norte;rio grande do sul;rondonia;roraima;santa catarina;sao paulo;sergipe;tocantins".split(";"),
))


def fold(value: Any) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(value or "").lower())
                   if not unicodedata.combining(c))


class InssBenefitsClient(StjClient):
    BASE = "https://dadosabertos.inss.gov.br/api/3/action"
    DATASET = "beneficios-indeferidos-plano-de-dados-abertos-jun-2023-a-jun-2025"

    def monthly_resource(self, package: dict[str, Any], competence: int) -> dict[str, Any]:
        year, month = divmod(competence, 100)
        if not 2000 <= year <= 2100 or not 1 <= month <= 12:
            raise ValueError("Competência INSS deve ser AAAAMM")
        resources = [r for r in package.get("resources", [])
                     if str(r.get("format", "")).upper() == "XLSX"
                     and re.search(rf"\b{year}\b", fold(r.get("name")))
                     and re.search(rf"\b{MONTHS[month-1]}\b", fold(r.get("name")))]
        if len(resources) != 1:
            raise ValueError(f"Esperado um recurso XLSX para {competence}; encontrados {len(resources)}")
        return resources[0]


def aggregate_rows(rows, uf: str, competence: int | None = None) -> tuple[Counter, dict[str, Any]]:
    columns = {}
    header = None
    for header_line in range(1, 21):
        header = next(rows, None)
        if header is None:
            break
        columns = {}
        for index, value in enumerate(header):
            key = re.sub(r"[^a-z0-9]", "", fold(value))
            columns.setdefault(key, []).append(index)
        if "especie" in columns and "uf" in columns:
            break
    if header is None:
        raise ValueError("Planilha INSS sem cabeçalho")
    required = {"especie": ("especie", "codigodaespecie"),
                "motivo": ("motivoindeferimento", "motivodoindeferimento", "motivo"),
                "uf": ("uf", "unidadefederacao", "unidadedafederacao")}
    indices = {}
    for name, alternatives in required.items():
        matches = [index for key in alternatives for index in columns.get(key, [])]
        # Layout real INSS: duas colunas Espécie consecutivas (código, descrição).
        # Selecionar o código e exigir valor numérico em TODAS as linhas.
        if name == "especie" and len(matches) == 2 and matches[1] == matches[0] + 1:
            matches = matches[:1]
        if len(matches) != 1:
            raise ValueError(f"Coluna INSS {name} ausente/ambígua: {list(columns)}")
        indices[name] = matches[0]
    if competence is not None:
        matches = columns.get("competenciaindeferimento", [])
        if len(matches) != 1:
            raise ValueError("Coluna competência de indeferimento ausente/ambígua")
        indices["competencia"] = matches[0]
    if uf.upper() not in UF_NAMES:
        raise ValueError("UF brasileira inválida")
    counts = Counter()
    read = selected = unknown_state = 0
    for row in rows:
        if not any(value is not None for value in row):
            continue
        read += 1
        if len(row) <= max(indices.values()):
            raise ValueError(f"Linha {read+1} tem menos colunas que o cabeçalho")
        if competence is not None:
            value = row[indices["competencia"]]
            period = value.strftime("%Y%m") if isinstance(value, datetime) else str(value).strip()
            if period != str(competence):
                raise ValueError(f"Competência divergente na linha {read+header_line}; arquivo não corresponde ao recurso")
        species_text = str(row[indices["especie"]] or "").strip()
        match = re.match(r"^(\d+)(?:\s*[-.]\s*.*)?$", species_text)
        if not match:
            raise ValueError(f"Espécie inválida na linha {read+1}; importação interrompida")
        species = int(match.group(1))
        if species not in (87, 88):
            continue
        state = fold(row[indices["uf"]]).strip()
        target = uf.lower()
        if state.upper() not in UF_NAMES and state not in UF_NAMES.values():
            unknown_state += 1
        if state != target and state != UF_NAMES[uf.upper()]:
            continue
        reason = str(row[indices["motivo"]] or "").strip() or "Não informado"
        counts[(uf.upper(), species, reason)] += 1
        selected += 1
    return counts, {"linhas_lidas": read, "linhas_selecionadas": selected,
                    "uf": uf.upper(), "especies": [87, 88], "cabecalho": list(header),
                    "linha_cabecalho": header_line, "bpc_uf_desconhecida": unknown_state}


def aggregate_xlsx(body: bytes, uf: str, competence: int | None = None) -> tuple[Counter, dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Planilha INSS com múltiplas abas exige seleção explícita")
        return aggregate_rows(iter(workbook.worksheets[0].iter_rows(values_only=True)), uf, competence)
    finally:
        workbook.close()


def store_aggregates(session, counts: Counter, resource: dict[str, Any], competence: int,
                     body: bytes, raw: Path, run: str) -> int:
    digest = hashlib.sha256(body).hexdigest()
    inserted = 0
    for (uf, species, reason), quantity in counts.items():
        key = dict(fonte="inss_indeferimentos", competencia=competence, uf=uf, especie=species,
                   motivo=reason, hash_arquivo=digest)
        existing = session.scalar(select(IndicadorInssIndeferimento).filter_by(**key))
        if existing is not None:
            if existing.quantidade != quantity:
                raise ValueError("Agregação divergente para os mesmos bytes; verifique parser")
            continue
        session.add(IndicadorInssIndeferimento(**key, quantidade=quantity,
            recurso_id=str(resource["id"]), recurso_url=resource["url"], arquivo_bruto=str(raw),
            coleta_id=uuid.UUID(run), coletado_em=datetime.now(timezone.utc)))
        inserted += 1
    return inserted
